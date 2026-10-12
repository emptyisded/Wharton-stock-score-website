"""
engine.py  -  the analysis engine (27 indicators, sector/market ranges, scoring, Excel export).

WHERE TO EDIT
  1. INDICATORS list   -> replace each `fn` with YOUR formula (marked  # FORMULA )
  2. ideal_ranges.csv  -> your ideal range per market / sector / indicator (auto-created on first run)
"""
import os, math, tempfile, datetime as dt
from dataclasses import dataclass
from typing import Callable, Optional

import numpy as np
import pandas as pd
import yfinance as yf

RANGES_CSV = "ideal_ranges.csv"

# ----------------------------------------------------------------------------
# 1. Timeframes  (price-based indicators use these; fundamentals use latest filings)
# ----------------------------------------------------------------------------
TIMEFRAMES = {
    # label: (history kwargs, bars per year used for annualising)
    "1 Day":   (dict(period="1d", interval="5m"), 78 * 252),
    "5 Days":  (dict(period="5d", interval="30m"), 13 * 252),
    "6 Months": (dict(period="6mo", interval="1d"), 252),
    "1 Year":  (dict(period="1y", interval="1d"), 252),
    "3 Years": (dict(period="3y", interval="1d"), 252),
    "4 Years": (dict(start=(dt.date.today() - dt.timedelta(days=365 * 4 + 1)).isoformat(), interval="1d"), 252),
}

RISK_FREE = {"India": 0.068, "Global": 0.043}   # used for Sharpe ratio; edit freely


# ----------------------------------------------------------------------------
# 2. Helpers
# ----------------------------------------------------------------------------
def num(x) -> Optional[float]:
    try:
        v = float(x)
        return v if math.isfinite(v) else None
    except (TypeError, ValueError):
        return None


def ratio(a, b, scale=1.0) -> Optional[float]:
    a, b = num(a), num(b)
    if a is None or b is None or b == 0:
        return None
    return a / b * scale


def pct(x) -> Optional[float]:
    """yfinance gives fractions (0.25) -> we show percent (25)."""
    v = num(x)
    return None if v is None else v * 100


@dataclass
class Ctx:
    """Everything an indicator formula can use."""
    info: dict
    income: pd.DataFrame          # annual income statement (rows = items, cols = dates)
    balance: pd.DataFrame         # annual balance sheet
    close: pd.Series              # closing prices for the selected timeframe
    bars_per_year: int
    market: str                   # "India" or "Global"


def stmt_row(df: pd.DataFrame, *names) -> Optional[float]:
    if df is None or df.empty:
        return None
    for n in names:
        if n in df.index:
            s = df.loc[n].dropna()
            if len(s):
                return num(s.iloc[0])
    return None


def returns(c: Ctx) -> pd.Series:
    return c.close.pct_change().dropna()


# ----------------------------------------------------------------------------
# 3. THE 27 INDICATORS  -  replace formulas with yours.  Each fn(ctx) -> float | None
#    Units: ratios as plain numbers, anything "%" is expressed as a percent (25 = 25%).
# ----------------------------------------------------------------------------
@dataclass
class Indicator:
    key: str
    name: str
    group: str
    unit: str
    formula: str                      # human-readable text shown in the Excel file
    fn: Callable[[Ctx], Optional[float]]
    default_range: tuple              # (low, high); None = open ended


def _roic(c: Ctx):
    ebit = stmt_row(c.income, "EBIT", "Operating Income")
    tax_rate = num(stmt_row(c.income, "Tax Rate For Calcs")) or 0.21
    invested = stmt_row(c.balance, "Invested Capital")
    if ebit is None or invested in (None, 0):
        return None
    return ebit * (1 - tax_rate) / invested * 100


def _interest_cov(c: Ctx):
    ebit = stmt_row(c.income, "EBIT", "Operating Income")
    interest = stmt_row(c.income, "Interest Expense")
    if ebit is None or not interest:
        return None
    return ebit / abs(interest)


def _fcf_yield(c: Ctx):
    return ratio(c.info.get("freeCashflow"), c.info.get("marketCap"), 100)


def _div_yield(c: Ctx):
    price = c.info.get("currentPrice") or c.info.get("regularMarketPrice")
    r = ratio(c.info.get("dividendRate"), price, 100)
    return 0.0 if r is None and c.info.get("dividendRate") is None else r


def _net_debt_ebitda(c: Ctx):
    debt, cash = num(c.info.get("totalDebt")), num(c.info.get("totalCash"))
    if debt is None or cash is None:
        return None
    return ratio(debt - cash, c.info.get("ebitda"))


def _total_return(c: Ctx):
    s = c.close.dropna()
    return None if len(s) < 2 else (s.iloc[-1] / s.iloc[0] - 1) * 100


def _volatility(c: Ctx):
    r = returns(c)
    return None if len(r) < 2 else r.std() * math.sqrt(c.bars_per_year) * 100


def _max_drawdown(c: Ctx):
    s = c.close.dropna()
    return None if len(s) < 2 else ((s / s.cummax()) - 1).min() * 100


def _rsi(c: Ctx, n=14):
    d = c.close.diff().dropna()
    if len(d) < n + 1:
        return None
    up, dn = d.clip(lower=0), (-d).clip(lower=0)
    rs = up.ewm(alpha=1 / n, adjust=False).mean().iloc[-1] / max(dn.ewm(alpha=1 / n, adjust=False).mean().iloc[-1], 1e-12)
    return 100 - 100 / (1 + rs)


def _sharpe(c: Ctx):
    r = returns(c)
    if len(r) < 5 or r.std() == 0:
        return None
    rf = RISK_FREE.get(c.market, 0.04) / c.bars_per_year
    return (r.mean() - rf) / r.std() * math.sqrt(c.bars_per_year)


INDICATORS = [
    # ---- Valuation -------------------------------------------------------
    Indicator("pe", "P/E (trailing)", "Valuation", "x", "Price / EPS (TTM)",
              lambda c: num(c.info.get("trailingPE")), (10, 25)),                       # FORMULA
    Indicator("fwd_pe", "Forward P/E", "Valuation", "x", "Price / forward EPS",
              lambda c: num(c.info.get("forwardPE")), (8, 22)),                         # FORMULA
    Indicator("pb", "Price / Book", "Valuation", "x", "Price / book value per share",
              lambda c: num(c.info.get("priceToBook")), (1, 4)),                        # FORMULA
    Indicator("ps", "Price / Sales", "Valuation", "x", "Market cap / revenue (TTM)",
              lambda c: num(c.info.get("priceToSalesTrailing12Months")), (1, 5)),       # FORMULA
    Indicator("ev_ebitda", "EV / EBITDA", "Valuation", "x", "Enterprise value / EBITDA",
              lambda c: num(c.info.get("enterpriseToEbitda")), (6, 15)),                # FORMULA
    Indicator("peg", "PEG ratio", "Valuation", "x", "P/E divided by earnings growth",
              lambda c: num(c.info.get("trailingPegRatio") or c.info.get("pegRatio")), (0.5, 1.5)),  # FORMULA
    # ---- Profitability ---------------------------------------------------
    Indicator("roe", "Return on equity", "Profitability", "%", "Net income / shareholder equity",
              lambda c: pct(c.info.get("returnOnEquity")), (15, None)),                 # FORMULA
    Indicator("roa", "Return on assets", "Profitability", "%", "Net income / total assets",
              lambda c: pct(c.info.get("returnOnAssets")), (5, None)),                  # FORMULA
    Indicator("roic", "Return on invested capital", "Profitability", "%", "EBIT x (1 - tax) / invested capital",
              _roic, (10, None)),                                                       # FORMULA
    Indicator("gross_m", "Gross margin", "Profitability", "%", "Gross profit / revenue",
              lambda c: pct(c.info.get("grossMargins")), (30, None)),                   # FORMULA
    Indicator("op_m", "Operating margin", "Profitability", "%", "Operating income / revenue",
              lambda c: pct(c.info.get("operatingMargins")), (15, None)),               # FORMULA
    Indicator("net_m", "Net margin", "Profitability", "%", "Net income / revenue",
              lambda c: pct(c.info.get("profitMargins")), (10, None)),                  # FORMULA
    # ---- Growth ----------------------------------------------------------
    Indicator("rev_g", "Revenue growth (YoY)", "Growth", "%", "Latest quarter revenue vs same quarter last year",
              lambda c: pct(c.info.get("revenueGrowth")), (5, None)),                   # FORMULA
    Indicator("earn_g", "Earnings growth (YoY)", "Growth", "%", "Latest quarter earnings vs same quarter last year",
              lambda c: pct(c.info.get("earningsGrowth")), (5, None)),                  # FORMULA
    Indicator("eps_g", "Quarterly earnings growth", "Growth", "%", "Quarterly net income growth",
              lambda c: pct(c.info.get("earningsQuarterlyGrowth")), (5, None)),         # FORMULA
    # ---- Leverage and liquidity -----------------------------------------
    Indicator("de", "Debt / Equity", "Financial health", "x", "Total debt / shareholder equity",
              lambda c: None if num(c.info.get("debtToEquity")) is None else num(c.info.get("debtToEquity")) / 100,
              (0, 1)),                                                                  # FORMULA
    Indicator("current", "Current ratio", "Financial health", "x", "Current assets / current liabilities",
              lambda c: num(c.info.get("currentRatio")), (1.2, 3)),                     # FORMULA
    Indicator("quick", "Quick ratio", "Financial health", "x", "(Cash + receivables) / current liabilities",
              lambda c: num(c.info.get("quickRatio")), (1, 2.5)),                       # FORMULA
    Indicator("int_cov", "Interest coverage", "Financial health", "x", "EBIT / interest expense",
              _interest_cov, (3, None)),                                                # FORMULA
    Indicator("nd_ebitda", "Net debt / EBITDA", "Financial health", "x", "(Total debt - cash) / EBITDA",
              _net_debt_ebitda, (0, 3)),                                                # FORMULA
    # ---- Cash returns ----------------------------------------------------
    Indicator("fcf_y", "Free cash flow yield", "Cash returns", "%", "Free cash flow / market cap",
              _fcf_yield, (3, None)),                                                   # FORMULA
    Indicator("div_y", "Dividend yield", "Cash returns", "%", "Annual dividend / price",
              _div_yield, (1, None)),                                                   # FORMULA
    # ---- Price behaviour (uses the selected timeframe) ------------------
    Indicator("ret", "Price return", "Price behaviour", "%", "Last close / first close - 1",
              _total_return, (0, None)),                                                # FORMULA
    Indicator("vol", "Volatility (annualised)", "Price behaviour", "%", "Std dev of returns x sqrt(bars/year)",
              _volatility, (0, 35)),                                                    # FORMULA
    Indicator("mdd", "Maximum drawdown", "Price behaviour", "%", "Worst peak-to-trough fall",
              _max_drawdown, (-25, 0)),                                                 # FORMULA
    Indicator("rsi", "RSI (14)", "Price behaviour", "", "Wilder RSI, 14 bars",
              _rsi, (40, 65)),                                                          # FORMULA
    Indicator("sharpe", "Sharpe ratio", "Price behaviour", "", "(Mean return - risk free) / std dev",
              _sharpe, (1, None)),                                                      # FORMULA
]
assert len(INDICATORS) == 27, f"Expected 27 indicators, found {len(INDICATORS)}"


# ----------------------------------------------------------------------------
# 4. Ideal ranges: CSV  (market, sector, indicator_key, low, high)
#    Lookup order: exact market+sector -> market + "*" -> built-in default
# ----------------------------------------------------------------------------
SECTORS = ["Technology", "Healthcare", "Financial Services", "Consumer Cyclical", "Consumer Defensive",
           "Industrials", "Energy", "Utilities", "Real Estate", "Basic Materials", "Communication Services"]


def create_template_csv():
    rows = []
    for m in ("Global", "India"):
        for s in SECTORS:
            for ind in INDICATORS:
                lo, hi = ind.default_range
                rows.append(dict(market=m, sector=s, indicator_key=ind.key, indicator=ind.name,
                                 low="" if lo is None else lo, high="" if hi is None else hi))
    pd.DataFrame(rows).to_csv(RANGES_CSV, index=False)


def load_ranges() -> dict:
    if not os.path.exists(RANGES_CSV):
        create_template_csv()
    df = pd.read_csv(RANGES_CSV)
    out = {}
    for r in df.itertuples():
        out[(str(r.market).strip(), str(r.sector).strip(), str(r.indicator_key).strip())] = (num(r.low), num(r.high))
    return out


def get_range(ranges, market, sector, ind: Indicator):
    for k in ((market, sector, ind.key), (market, "*", ind.key)):
        if k in ranges:
            return ranges[k]
    return ind.default_range


# ----------------------------------------------------------------------------
# 5. Scoring:  inside range = 100; outside decays linearly with distance
# ----------------------------------------------------------------------------
def score_value(v, lo, hi):
    if v is None:
        return None, "No data"
    if lo is not None and v < lo:
        dist, status, ref = lo - v, "Below", lo
    elif hi is not None and v > hi:
        dist, status, ref = v - hi, "Above", hi
    else:
        return 100.0, "Within"
    width = (hi - lo) if (lo is not None and hi is not None) else abs(ref)
    scale = max(abs(width), abs(ref) * 0.5, 1e-9)
    return max(0.0, 100 - 100 * dist / scale), status


def rating(s):
    if s >= 80: return "Excellent"
    if s >= 65: return "Strong"
    if s >= 50: return "Average"
    if s >= 35: return "Weak"
    return "Poor"


# ----------------------------------------------------------------------------
# 6. Data + analysis
# ----------------------------------------------------------------------------
def resolve_ticker(text: str):
    t = text.strip().upper().replace(" ", "")
    if not t:
        raise ValueError("Enter a ticker, for example NVDA or RELIANCE.NS")
    candidates = [t] if "." in t else [t, t + ".NS", t + ".BO"]
    for cand in candidates:
        try:
            h = yf.Ticker(cand).history(period="5d")
            if not h.empty:
                return cand
        except Exception:
            pass
    try:                                           # fall back to name search ("nvidia")
        for q in yf.Search(text.strip(), max_results=5).quotes:
            if q.get("symbol") and q.get("quoteType") in ("EQUITY", "ETF"):
                return q["symbol"]
    except Exception:
        pass
    raise ValueError(f"Could not find '{text}'. Try the exact Yahoo ticker (e.g. TCS.NS for India).")


def detect_market(symbol: str, info: dict) -> str:
    if symbol.endswith((".NS", ".BO")) or info.get("country") == "India":
        return "India"
    return "Global"


def analyze(user_text: str, timeframe: str):
    symbol = resolve_ticker(user_text)
    tk = yf.Ticker(symbol)
    info = tk.info or {}
    market = detect_market(symbol, info)
    sector = info.get("sector") or "Unknown"

    kwargs, bpy = TIMEFRAMES[timeframe]
    hist = tk.history(auto_adjust=True, **kwargs)
    if hist.empty:
        raise ValueError("No price history returned for this timeframe.")
    try:
        income, balance = tk.income_stmt, tk.balance_sheet
    except Exception:
        income, balance = pd.DataFrame(), pd.DataFrame()

    ctx = Ctx(info, income, balance, hist["Close"], bpy, market)
    ranges = load_ranges()

    rows = []
    for ind in INDICATORS:
        try:
            v = ind.fn(ctx)
        except Exception:
            v = None
        lo, hi = get_range(ranges, market, sector, ind)
        s, status = score_value(v, lo, hi)
        rows.append(dict(Group=ind.group, Indicator=ind.name, Value=v, Unit=ind.unit, Low=lo, High=hi,
                         Status=status, Score=s, Formula=ind.formula))
    df = pd.DataFrame(rows)

    valid = df["Score"].dropna()
    overall = float(valid.mean()) if len(valid) else 0.0
    meta = dict(symbol=symbol, name=info.get("longName") or info.get("shortName") or symbol, sector=sector,
                industry=info.get("industry", "-"), market=market, currency=info.get("currency", ""),
                timeframe=timeframe, price=num(hist["Close"].iloc[-1]), overall=overall, rating=rating(overall),
                coverage=f"{len(valid)} of {len(INDICATORS)}",
                generated=dt.datetime.now().strftime("%d %b %Y, %H:%M"))
    return meta, df


def fmt_range(lo, hi, unit):
    u = unit if unit in ("%", "x") else ""
    if lo is None and hi is None: return "Any"
    if lo is None: return f"up to {hi:g}{u}"
    if hi is None: return f"{lo:g}{u} and above"
    return f"{lo:g}{u} to {hi:g}{u}"


def fmt_val(v, unit):
    if v is None: return "N/A"
    u = unit if unit in ("%", "x") else ""
    return f"{v:,.2f}{u}"


# ----------------------------------------------------------------------------
# 7. Excel export
# ----------------------------------------------------------------------------
def to_excel(meta: dict, df: pd.DataFrame) -> str:
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    ink, paper = "14213D", "F4F6F9"
    head_fill = PatternFill("solid", fgColor=ink)
    fills = {"Within": "D7F0E3", "Below": "FBE3C8", "Above": "FBE3C8", "No data": "E8EAED"}
    thin = Side(style="thin", color="D5D9E0")

    ws = wb.active
    ws.title = "Summary"
    ws["A1"] = f"{meta['name']} ({meta['symbol']})"
    ws["A1"].font = Font(size=18, bold=True, color=ink)
    items = [("Overall score", f"{meta['overall']:.1f} / 100"), ("Rating", meta["rating"]),
             ("Sector", meta["sector"]), ("Industry", meta["industry"]), ("Market", meta["market"]),
             ("Timeframe", meta["timeframe"]), ("Last price", f"{meta['price']:,.2f} {meta['currency']}"),
             ("Indicators scored", meta["coverage"]), ("Generated", meta["generated"])]
    for i, (k, v) in enumerate(items, start=3):
        ws.cell(i, 1, k).font = Font(bold=True)
        ws.cell(i, 2, v)
    r0 = len(items) + 5
    ws.cell(r0, 1, "Score by group").font = Font(bold=True, size=12, color=ink)
    g = df.groupby("Group", sort=False)["Score"].mean()
    for i, (name, val) in enumerate(g.items(), start=r0 + 1):
        ws.cell(i, 1, name)
        ws.cell(i, 2, None if pd.isna(val) else round(float(val), 1))
    ws.column_dimensions["A"].width = 26
    ws.column_dimensions["B"].width = 28

    ws2 = wb.create_sheet("Indicators")
    heads = ["#", "Group", "Indicator", "Value", "Ideal range", "Status", "Score (0-100)", "Formula"]
    for c, h in enumerate(heads, 1):
        cell = ws2.cell(1, c, h)
        cell.font, cell.fill = Font(bold=True, color="FFFFFF"), head_fill
        cell.alignment = Alignment(horizontal="center", vertical="center")
    for i, r in enumerate(df.itertuples(), start=1):
        vals = [i, r.Group, r.Indicator, fmt_val(r.Value, r.Unit), fmt_range(r.Low, r.High, r.Unit),
                r.Status, None if pd.isna(r.Score) else round(r.Score, 1), r.Formula]
        for c, v in enumerate(vals, 1):
            cell = ws2.cell(i + 1, c, v)
            cell.border = Border(top=thin, bottom=thin, left=thin, right=thin)
            if c == 6:
                cell.fill = PatternFill("solid", fgColor=fills.get(r.Status, paper))
            if c in (1, 4, 5, 6, 7):
                cell.alignment = Alignment(horizontal="center")
    for c, w in enumerate([5, 18, 30, 14, 22, 12, 14, 48], 1):
        ws2.column_dimensions[get_column_letter(c)].width = w
    ws2.freeze_panes = "A2"

    path = os.path.join(tempfile.gettempdir(), f"{meta['symbol']}_{meta['timeframe'].replace(' ', '')}_score.xlsx")
    wb.save(path)
    return path


